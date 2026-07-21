"""
Indexation des documents pédagogiques depuis /Users/danielmorel/
vers Supabase — Daniel Morel Éternel
"""

import hashlib
import os
import sys
from pathlib import Path

import psycopg2
from dotenv import load_dotenv
from sentence_transformers import SentenceTransformer

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from pipeline.paths import ROOT

load_dotenv(ROOT / ".env")

DB_URL = os.getenv("SUPABASE_DB_URL", "")
CHUNK_SIZE = 800
OVERLAP = 150

DOSSIERS = [
    "/Users/danielmorel/Documents",
    "/Users/danielmorel/Desktop",
    "/Users/danielmorel/Downloads",
]

MOTS_CLES = ["pnl","ennea","hypnos","formation","pedago","conference","praticien",
              "coaching","management","cohesion","entreprise","programme","synopsis",
              "base","module","leadership","communication","mindset","ego","reconnexion"]

EXCLURE = ["facture","convention","presence","attestation","comptab","banque",
           "assur","feuille","devis","contrat","rib","kbis","logo","signature"]

EXTENSIONS = {".pdf", ".docx", ".pptx", ".doc"}

def extraire_texte(path):
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            r = PdfReader(str(path))
            return " ".join(p.extract_text() or "" for p in r.pages)
        elif ext in (".docx", ".doc"):
            from docx import Document
            doc = Document(str(path))
            return " ".join(p.text for p in doc.paragraphs)
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(str(path))
            textes = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        textes.append(shape.text)
            return " ".join(textes)
    except Exception:
        return ""
    return ""

def pertinent(path):
    nom = path.name.lower()
    if not any(nom.endswith(e) for e in EXTENSIONS):
        return False
    if any(x in nom for x in EXCLURE):
        return False
    if any(x in nom for x in MOTS_CLES):
        return True
    # Vérifier le dossier parent
    chemin = str(path).lower()
    return any(x in chemin for x in MOTS_CLES)

def chunker(texte):
    chunks, i = [], 0
    while i < len(texte):
        chunks.append(texte[i:i+CHUNK_SIZE])
        i += CHUNK_SIZE - OVERLAP
    return chunks

def hash_texte(t):
    return hashlib.md5(t.encode()).hexdigest()

print("Chargement du modèle d'embeddings...")
model = SentenceTransformer("paraphrase-multilingual-MiniLM-L12-v2")

print("Connexion Supabase...")
conn = psycopg2.connect(DB_URL)
cur = conn.cursor()

# Récupérer les hash déjà présents pour éviter les doublons
cur.execute("SELECT md5(contenu) FROM documents")
deja_indexes = set(r[0] for r in cur.fetchall())
print(f"{len(deja_indexes)} chunks déjà dans la base")

# Collecter les fichiers
fichiers = []
for dossier in DOSSIERS:
    for path in Path(dossier).rglob("*"):
        if ".git" in str(path) or "miniconda3" in str(path) or "Library" in str(path):
            continue
        if path.is_file() and pertinent(path):
            fichiers.append(path)

print(f"{len(fichiers)} fichiers à traiter\n")

total_chunks = 0
erreurs = 0

for i, path in enumerate(fichiers, 1):
    try:
        texte = extraire_texte(path).strip()
        if len(texte) < 100:
            continue

        chunks = chunker(texte)
        nouveaux = [c for c in chunks if hash_texte(c) not in deja_indexes]
        if not nouveaux:
            continue

        source = str(path).replace("/Users/danielmorel/", "")
        embeddings = model.encode(nouveaux)

        for chunk, emb in zip(nouveaux, embeddings):
            cur.execute(
                "INSERT INTO documents (contenu, source, embedding) VALUES (%s, %s, %s)",
                (chunk, f"Documents_Mac/{source}", str(emb.tolist()))
            )
            deja_indexes.add(hash_texte(chunk))

        conn.commit()
        total_chunks += len(nouveaux)
        print(f"[{i}/{len(fichiers)}] ✅ {path.name} — {len(nouveaux)} chunks")

    except Exception as e:
        erreurs += 1
        print(f"[{i}/{len(fichiers)}] ⚠️  {path.name} — {e}")

cur.close()
conn.close()
print(f"\n{'='*50}")
print(f"Terminé — {total_chunks} chunks ajoutés, {erreurs} erreurs")
