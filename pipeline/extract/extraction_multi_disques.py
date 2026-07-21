"""
Extraction multi-disques + déduplication par contenu — Daniel Morel Éternel
Parcourt tous les disques, extrait le texte, supprime les doublons par hash.
"""

import hashlib
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from pipeline.paths import LOGS_DIR, TEXTES_DIR

SOURCES = [
    Path("/Volumes/Securité 500 g"),
    Path("/Volumes/clean Disk"),
    Path("/Volumes/hysta"),
]

DEST = TEXTES_DIR
DEST.mkdir(parents=True, exist_ok=True)

LOG_OK = LOGS_DIR / "extraction_ok.log"
LOG_ERR = LOGS_DIR / "extraction_erreurs.log"
LOG_DUP = LOGS_DIR / "extraction_doublons.log"

EXTENSIONS = {".docx", ".pdf", ".pptx", ".ppt", ".doc"}

# Charger les hashs déjà indexés (textes extraits de Cle 2T)
def charger_hashs_existants():
    hashs = set()
    for fichier_txt in DEST.rglob("*.txt"):
        try:
            contenu = fichier_txt.read_text(encoding="utf-8", errors="ignore")
            hashs.add(hashlib.md5(contenu.encode()).hexdigest())
        except Exception:
            pass
    return hashs


def extraire_docx(path):
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extraire_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    pages = []
    for page in reader.pages:
        texte = page.extract_text()
        if texte:
            pages.append(texte)
    return "\n".join(pages)


def extraire_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    lignes = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                lignes.append(shape.text)
    return "\n".join(lignes)


def traiter(path):
    ext = path.suffix.lower()
    if ext in (".docx", ".doc"):
        return extraire_docx(path)
    elif ext == ".pdf":
        return extraire_pdf(path)
    elif ext in (".pptx", ".ppt"):
        return extraire_pptx(path)
    return None


def nom_dest(source_path, source_root):
    try:
        relative = source_path.relative_to(source_root)
    except ValueError:
        relative = Path(source_path.name)
    # Préfixer par le nom du disque pour éviter les conflits de noms
    disque = source_root.name.replace(" ", "_")
    dest = DEST / disque / relative.with_suffix(".txt")
    dest.parent.mkdir(parents=True, exist_ok=True)
    return dest


def main():
    print("Chargement des hashs existants (Cle 2T)...")
    hashs_vus = charger_hashs_existants()
    print(f"{len(hashs_vus)} textes déjà indexés depuis Cle 2T\n")

    total_ok = 0
    total_dup = 0
    total_err = 0

    with open(LOG_OK, "a", encoding="utf-8") as fok, \
         open(LOG_ERR, "a", encoding="utf-8") as ferr, \
         open(LOG_DUP, "w", encoding="utf-8") as fdup:

        for source in SOURCES:
            if not source.exists():
                print(f"Disque non trouvé : {source}")
                continue

            fichiers = [
                p for p in source.rglob("*")
                if p.is_file() and p.suffix.lower() in EXTENSIONS
            ]
            print(f"=== {source.name} — {len(fichiers)} fichiers ===")

            for i, fichier in enumerate(fichiers, 1):
                if i % 200 == 0:
                    print(f"  {i}/{len(fichiers)} — {fichier.name}")
                try:
                    texte = traiter(fichier)
                    if not texte or not texte.strip():
                        ferr.write(f"VIDE: {fichier}\n")
                        total_err += 1
                        continue

                    h = hashlib.md5(texte.encode()).hexdigest()
                    if h in hashs_vus:
                        fdup.write(f"DOUBLON: {fichier}\n")
                        total_dup += 1
                        continue

                    hashs_vus.add(h)
                    dest = nom_dest(fichier, source)
                    dest.write_text(texte, encoding="utf-8")
                    fok.write(str(fichier) + "\n")
                    total_ok += 1

                except Exception as e:
                    ferr.write(f"ERREUR: {fichier} — {e}\n")
                    total_err += 1

            print(f"  → {source.name} terminé\n")

    print(f"\n{'='*50}")
    print(f"Nouveaux extraits uniques : {total_ok}")
    print(f"Doublons ignorés          : {total_dup}")
    print(f"Erreurs/vides             : {total_err}")
    print(f"Textes dans : {DEST}")


if __name__ == "__main__":
    main()
