"""
Extraction de texte — Daniel Morel Éternel
Parcourt le disque Cle 2T et convertit .docx / .pdf / .pptx en .txt
"""

import os
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from pipeline.paths import LOGS_DIR, TEXTES_DIR

SOURCE = Path("/Volumes/Cle 2T")
DEST = TEXTES_DIR
DEST.mkdir(parents=True, exist_ok=True)

LOG_OK = LOGS_DIR / "extraction_ok.log"
LOG_ERR = LOGS_DIR / "extraction_erreurs.log"

EXTENSIONS = {".docx", ".pdf", ".pptx", ".ppt", ".doc"}


def extraire_docx(path):
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extraire_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    pages = []
    for page in reader.pages:
        texte = page.extract_text()
        if texte:
            pages.append(texte)
    return "\n".join(pages)


def extraire_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    lignes = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                lignes.append(shape.text)
    return "\n".join(lignes)


def nom_dest(source_path):
    # Reproduit l'arborescence relative dans le dossier destination
    try:
        relative = source_path.relative_to(SOURCE)
    except ValueError:
        relative = Path(source_path.name)
    dest = DEST / relative.with_suffix(".txt")
    dest.parent.mkdir(parents=True, exist_ok=True)
    return dest


def traiter(path):
    ext = path.suffix.lower()
    if ext in (".docx", ".doc"):
        return extraire_docx(path)
    elif ext == ".pdf":
        return extraire_pdf(path)
    elif ext in (".pptx", ".ppt"):
        return extraire_pptx(path)
    return None


def main():
    fichiers = [
        p for p in SOURCE.rglob("*")
        if p.is_file() and p.suffix.lower() in EXTENSIONS
    ]
    total = len(fichiers)
    print(f"{total} fichiers trouvés. Démarrage de l'extraction...\n")

    ok = 0
    erreurs = 0

    with open(LOG_OK, "w", encoding="utf-8") as fok, \
         open(LOG_ERR, "w", encoding="utf-8") as ferr:

        for i, fichier in enumerate(fichiers, 1):
            if i % 100 == 0 or i == 1:
                print(f"  {i}/{total} — {fichier.name}")
            try:
                texte = traiter(fichier)
                if texte and texte.strip():
                    dest = nom_dest(fichier)
                    dest.write_text(texte, encoding="utf-8")
                    fok.write(str(fichier) + "\n")
                    ok += 1
                else:
                    ferr.write(f"VIDE: {fichier}\n")
                    erreurs += 1
            except Exception as e:
                ferr.write(f"ERREUR: {fichier} — {e}\n")
                erreurs += 1

    print(f"\nTerminé — {ok} extraits avec succès, {erreurs} erreurs/vides")
    print(f"Textes dans : {DEST}")
    print(f"Log erreurs : {LOG_ERR}")


if __name__ == "__main__":
    main()
